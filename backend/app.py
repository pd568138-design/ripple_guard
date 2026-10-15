"""RippleGuard API - supply-chain disruption and recovery prototype.

For classroom/demo use. Set MONGO_URI and SECRET_KEY in environment for deployment.
OTP is simulated unless a real SMS provider is integrated.
"""
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from pymongo import MongoClient
from bson import ObjectId
from datetime import datetime, timedelta, timezone
from functools import wraps
import os, secrets, hashlib, io, csv, json
import jwt
from werkzeug.security import generate_password_hash, check_password_hash
from dotenv import load_dotenv

load_dotenv()

app = Flask(__name__)
CORS(app, resources={r"/api/*": {"origins": os.getenv("CORS_ORIGINS", "*").split(",")}})
SECRET_KEY = os.getenv("SECRET_KEY", "change-this-development-secret-before-deploying")
OTP_TTL_SECONDS = int(os.getenv("OTP_TTL_SECONDS", "180"))
DEMO_OTP = os.getenv("DEMO_OTP", "true").lower() == "true"
MONGO_URI = os.getenv("MONGO_URI", "")

client = None
mongo_db = None
if MONGO_URI:
    try:
        _kw = {"serverSelectionTimeoutMS": 8000}
        if "mongodb+srv" in MONGO_URI or "mongodb.net" in MONGO_URI:
            try:
                import certifi
                _kw["tlsCAFile"] = certifi.where()
            except ImportError:
                pass
        client = MongoClient(MONGO_URI, **_kw)
        client.admin.command("ping")
        mongo_db = client[os.getenv("MONGO_DB", "rippleguard")]
    except Exception as exc:
        print("MongoDB connection unavailable; running in local demo mode:", exc)
        mongo_db = None

# A simple in-memory fallback keeps the app runnable before Atlas is configured.
_memory = {name: [] for name in ["users", "suppliers", "facilities", "products", "orders", "disruptions", "notifications", "analyses", "pending_registrations", "process_steps"]}


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def clean_doc(doc):
    if not doc:
        return doc
    out = dict(doc)
    out["id"] = str(out.pop("_id", out.get("id", "")))
    return out


def col(name):
    return mongo_db[name] if mongo_db is not None else None


def find_docs(name, query=None, limit=1000):
    query = query or {}
    if col(name) is not None:
        return [clean_doc(x) for x in col(name).find(query).sort("created_at", -1).limit(limit)]
    return [dict(x) for x in _memory[name] if all(x.get(k) == v for k, v in query.items())][:limit]


def find_one(name, query):
    if col(name) is not None:
        x = col(name).find_one(query)
        return clean_doc(x) if x else None
    for x in _memory[name]:
        if all(x.get(k) == v for k, v in query.items()):
            return dict(x)
    return None


def insert_doc(name, doc):
    doc = dict(doc)
    doc.setdefault("created_at", now_iso())
    if col(name) is not None:
        result = col(name).insert_one(doc)
        doc["id"] = str(result.inserted_id)
        doc.pop("_id", None)
    else:
        doc["id"] = secrets.token_hex(12)
        _memory[name].insert(0, doc)
    return doc


def update_doc(name, doc_id, updates, owner_id=None, workspace=None):
    if col(name) is not None:
        q = {"_id": ObjectId(doc_id)}
        if owner_id: q["owner_id"] = owner_id
        if workspace: q["workspace"] = workspace
        result = col(name).update_one(q, {"$set": updates})
        if result.matched_count:
            x = col(name).find_one(q)
            return clean_doc(x)
        return None
    for x in _memory[name]:
        if x.get("id") == doc_id and (not owner_id or x.get("owner_id") == owner_id) and (not workspace or x.get("workspace") == workspace):
            x.update(updates)
            return dict(x)
    return None


def delete_doc(name, doc_id, owner_id=None, workspace=None):
    if col(name) is not None:
        q = {"_id": ObjectId(doc_id)}
        if owner_id: q["owner_id"] = owner_id
        if workspace: q["workspace"] = workspace
        return col(name).delete_one(q).deleted_count > 0
    for i, x in enumerate(_memory[name]):
        if x.get("id") == doc_id and (not owner_id or x.get("owner_id") == owner_id) and (not workspace or x.get("workspace") == workspace):
            _memory[name].pop(i)
            return True
    return False


def workspace_of(user):
    return user.get("workspace") or str(user.get("company", "")).strip().lower() or user.get("id", "")


def owner_query(user):
    """Shared by everybody in the same company/workspace (all roles see the same records)."""
    return {"workspace": workspace_of(user)}


def user_query(user):
    """Private to one user (notifications)."""
    return {"owner_id": user["id"]}


def auth_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        header = request.headers.get("Authorization", "")
        if not header.startswith("Bearer "):
            return jsonify({"error": "Please log in to continue."}), 401
        try:
            payload = jwt.decode(header.split(" ", 1)[1], SECRET_KEY, algorithms=["HS256"])
            user = find_one("users", {"id": payload.get("sub")})
            if not user:
                # MongoDB users are keyed by ObjectId; fallback to email lookup below.
                user = find_one("users", {"email": payload.get("email")})
            if not user:
                return jsonify({"error": "Session is invalid. Please log in again."}), 401
            request.current_user = user
        except jwt.ExpiredSignatureError:
            return jsonify({"error": "Session expired. Please log in again."}), 401
        except Exception:
            return jsonify({"error": "Invalid authentication token."}), 401
        return fn(*args, **kwargs)
    return wrapped


def require_roles(*roles):
    def decorator(fn):
        @wraps(fn)
        @auth_required
        def wrapped(*args, **kwargs):
            if request.current_user.get("role") not in roles and request.current_user.get("role") != "Admin":
                return jsonify({"error": "Your role does not have permission for this action."}), 403
            return fn(*args, **kwargs)
        return wrapped
    return decorator


def add_notification(user_id, title, message, level="info", actor_id=None):
    return insert_doc("notifications", {"owner_id": user_id, "title": title, "message": message, "level": level, "read": False, "actor": actor_id, "created_at": now_iso()})


def notify_workspace(user, title, message, level="info"):
    ws = workspace_of(user)
    for u in find_docs("users", {}):
        if workspace_of(u) == ws:
            add_notification(u["id"], title, message, level, actor_id=user["id"])


@app.get("/api/health")
def health():
    return jsonify({"status": "ok", "service": "RippleGuard API", "database": "MongoDB Atlas" if mongo_db is not None else "Local demo memory (configure MONGO_URI for persistence)", "time": now_iso()})


@app.post("/api/auth/register")
def register():
    data = request.get_json(silent=True) or {}
    name = str(data.get("name", "")).strip()
    email = str(data.get("email", "")).strip().lower()
    phone = str(data.get("phone", "")).strip()
    password = str(data.get("password", ""))
    role = str(data.get("role", "Supply Chain Manager"))
    company = str(data.get("company", "")).strip()
    allowed_roles = ["Admin", "Supply Chain Manager", "Procurement Officer", "Warehouse Manager", "Risk Analyst", "Supplier"]
    if not all([name, email, phone, password, company]):
        return jsonify({"error": "Please complete all required fields."}), 400
    if "@" not in email or len(password) < 8:
        return jsonify({"error": "Enter a valid email and a password with at least 8 characters."}), 400
    if role not in allowed_roles:
        return jsonify({"error": "Choose a valid workspace role."}), 400
    if find_one("users", {"email": email}) or find_one("pending_registrations", {"email": email}):
        return jsonify({"error": "This email is already registered or awaiting verification."}), 409
    otp = f"{secrets.randbelow(1000000):06d}"
    expires = datetime.now(timezone.utc) + timedelta(seconds=OTP_TTL_SECONDS)
    pending = {"name": name, "email": email, "phone": phone, "company": company, "role": role,
               "password_hash": generate_password_hash(password), "otp_hash": hashlib.sha256(otp.encode()).hexdigest(),
               "expires_at": expires.isoformat(), "created_at": now_iso()}
    insert_doc("pending_registrations", pending)
    response = {"message": "Verification code created.", "expires_in": OTP_TTL_SECONDS, "delivery": "demo" if DEMO_OTP else "provider-not-configured"}
    if DEMO_OTP:
        response["demo_otp"] = otp
        response["notice"] = "Demo mode: code shown here instead of SMS. Connect an SMS provider before production use."
    return jsonify(response), 201


@app.post("/api/auth/verify-otp")
def verify_otp():
    data = request.get_json(silent=True) or {}
    email = str(data.get("email", "")).strip().lower()
    otp = str(data.get("otp", "")).strip()
    pending = find_one("pending_registrations", {"email": email})
    if not pending:
        return jsonify({"error": "Registration request not found. Please register again."}), 404
    try:
        expires = datetime.fromisoformat(pending["expires_at"])
        if datetime.now(timezone.utc) > expires:
            delete_doc("pending_registrations", pending["id"])
            return jsonify({"error": "OTP expired. Please register again."}), 410
    except Exception:
        return jsonify({"error": "OTP expiry could not be verified. Please register again."}), 400
    if not secrets.compare_digest(hashlib.sha256(otp.encode()).hexdigest(), pending["otp_hash"]):
        return jsonify({"error": "Incorrect OTP. Please check the code and try again."}), 400
    user = insert_doc("users", {"name": pending["name"], "email": pending["email"], "phone": pending["phone"],
                                 "company": pending["company"], "workspace": pending["company"].strip().lower(), "role": pending["role"], "password_hash": pending["password_hash"],
                                 "created_at": now_iso()})
    delete_doc("pending_registrations", pending["id"])
    add_notification(user["id"], "Welcome to RippleGuard", "Your workspace is ready. Add suppliers and facilities to begin.", "success")
    return jsonify({"message": "Phone/email verification completed. You can now log in."}), 200


@app.post("/api/auth/login")
def login():
    data = request.get_json(silent=True) or {}
    identifier = str(data.get("identifier", data.get("email", ""))).strip()
    password = str(data.get("password", ""))
    user = find_one("users", {"email": identifier.lower()})
    if not user:
        user = find_one("users", {"name": identifier})
    if not user or not check_password_hash(user.get("password_hash", ""), password):
        return jsonify({"error": "Incorrect email or password."}), 401
    token = jwt.encode({"sub": user["id"], "email": user["email"], "exp": datetime.now(timezone.utc) + timedelta(hours=8)}, SECRET_KEY, algorithm="HS256")
    safe_user = {k: user.get(k) for k in ["id", "name", "email", "phone", "company", "role"]}
    return jsonify({"token": token, "user": safe_user, "expires_in": 28800}), 200


@app.get("/api/auth/me")
@auth_required
def me():
    u = request.current_user
    return jsonify({"user": {k: u.get(k) for k in ["id", "name", "email", "phone", "company", "role"]}})


COLLECTIONS = {"suppliers", "facilities", "products", "orders", "disruptions", "process_steps"}

def can_write_collection(role, collection):
    if role in {"Admin", "Supply Chain Manager"}: return True
    if role == "Procurement Officer": return collection in {"suppliers", "products", "orders", "process_steps"}
    if role == "Warehouse Manager": return collection in {"facilities", "products", "orders", "disruptions", "process_steps"}
    if role == "Supplier": return collection in {"disruptions", "orders"}
    return False

@app.get("/api/<collection>")
@auth_required
def list_collection(collection):
    if collection not in COLLECTIONS: return jsonify({"error": "Unknown collection."}), 404
    items = find_docs(collection, owner_query(request.current_user))
    return jsonify({"items": items})

@app.post("/api/<collection>")
@auth_required
def create_collection(collection):
    if collection not in COLLECTIONS: return jsonify({"error": "Unknown collection."}), 404
    if not can_write_collection(request.current_user.get("role", ""), collection):
        return jsonify({"error": "Your role can view this area but cannot change these records."}), 403
    data = request.get_json(silent=True) or {}
    if not data.get("name") and collection in {"suppliers", "facilities", "products", "disruptions", "process_steps"}:
        return jsonify({"error": "A name is required."}), 400
    data.pop("id", None); data.pop("_id", None)
    data["owner_id"] = request.current_user["id"]
    data["workspace"] = workspace_of(request.current_user)
    data["created_at"] = now_iso()
    if collection == "suppliers":
        data.setdefault("capacity", 0); data.setdefault("unit_cost", 0); data.setdefault("lead_days", 1); data.setdefault("reliability", 95)
    if collection == "facilities":
        data.setdefault("type", "Factory"); data.setdefault("latitude", 11.0168); data.setdefault("longitude", 76.9558)
    if collection == "products":
        data.setdefault("required_components", 0); data.setdefault("priority", "Medium")
    if collection == "orders":
        data.setdefault("quantity", 0); data.setdefault("status", "Open"); data.setdefault("stages_completed", 0)
    if collection == "process_steps":
        data.setdefault("status", "In progress"); data.setdefault("started_at", now_iso())
    item = insert_doc(collection, data)
    _u = request.current_user
    notify_workspace(_u, f"New {collection[:-1] if collection.endswith('s') else collection} added", f"{item.get('name') or item.get('order_no') or 'Record'} was added by {_u.get('name')} ({_u.get('role')}).", "success")
    return jsonify({"item": item}), 201

@app.patch("/api/<collection>/<doc_id>")
@auth_required
def patch_collection(collection, doc_id):
    if collection not in COLLECTIONS: return jsonify({"error": "Unknown collection."}), 404
    if not can_write_collection(request.current_user.get("role", ""), collection):
        return jsonify({"error": "Your role cannot edit these records."}), 403
    data = request.get_json(silent=True) or {}
    for key in ["_id", "id", "owner_id", "workspace", "created_at", "password_hash"]: data.pop(key, None)
    item = update_doc(collection, doc_id, data, workspace=workspace_of(request.current_user))
    if not item: return jsonify({"error": "Record not found."}), 404
    _u = request.current_user
    notify_workspace(_u, f"{collection[:-1] if collection.endswith('s') else collection} updated", f"{item.get('name') or item.get('order_no') or 'Record'} was updated by {_u.get('name')} ({_u.get('role')}).", "info")
    return jsonify({"item": item})

@app.delete("/api/<collection>/<doc_id>")
@auth_required
def remove_collection(collection, doc_id):
    if collection not in COLLECTIONS: return jsonify({"error": "Unknown collection."}), 404
    if not can_write_collection(request.current_user.get("role", ""), collection):
        return jsonify({"error": "Your role cannot delete these records."}), 403
    if not delete_doc(collection, doc_id, workspace=workspace_of(request.current_user)): return jsonify({"error": "Record not found."}), 404
    return jsonify({"message": "Record deleted."})


@app.get("/api/notifications")
@auth_required
def notifications():
    items = find_docs("notifications", user_query(request.current_user), 100)
    unread = sum(1 for x in items if not x.get("read"))
    return jsonify({"items": items, "unread": unread})

@app.patch("/api/notifications/<doc_id>/read")
@auth_required
def mark_notification_read(doc_id):
    item = update_doc("notifications", doc_id, {"read": True}, request.current_user["id"])
    if not item: return jsonify({"error": "Notification not found."}), 404
    return jsonify({"item": item})

@app.post("/api/notifications/read-all")
@auth_required
def mark_all_read():
    items = find_docs("notifications", user_query(request.current_user), 1000)
    for item in items:
        if not item.get("read"):
            update_doc("notifications", item["id"], {"read": True}, request.current_user["id"])
    return jsonify({"message": "All notifications marked as read."})


def number(item, key, default=0):
    try: return float(item.get(key, default) or default)
    except (TypeError, ValueError): return float(default)


def build_analysis(user):
    suppliers = find_docs("suppliers", owner_query(user))
    facilities = find_docs("facilities", owner_query(user))
    products = find_docs("products", owner_query(user))
    orders = find_docs("orders", owner_query(user))
    disruptions = find_docs("disruptions", owner_query(user))
    total_required = sum(number(p, "required_components") for p in products)
    total_stock = sum(number(p, "available_stock") for p in products)
    shortage = max(0, total_required - total_stock)
    active_disruptions = [d for d in disruptions if d.get("status", "Active") == "Active"]
    impacted = []
    for s in suppliers:
        sname = str(s.get("name", "")).lower()
        matched = [d for d in active_disruptions if sname and sname in str(d.get("affected_supplier", "")).lower()]
        if matched:
            impacted.append({"supplier": s.get("name"), "supplier_id": s.get("id"), "events": [x.get("name") for x in matched], "capacity": number(s,"capacity"), "reliability": number(s,"reliability",95)})
    # A transparent heuristic score, not a trained predictive model.
    risk = min(100, round((len(active_disruptions)*18) + (shortage/max(1,total_required)*55) + (len(impacted)*12)))
    level = "Critical" if risk >= 70 else "Warning" if risk >= 35 else "Stable"
    return {"supplier_count": len(suppliers), "facility_count": len(facilities), "product_count": len(products), "order_count": len(orders),
            "active_disruptions": len(active_disruptions), "total_required": total_required, "total_stock": total_stock,
            "shortage": shortage, "risk_score": risk, "risk_level": level, "impacted_suppliers": impacted,
            "disruptions": disruptions, "suppliers": suppliers, "facilities": facilities, "products": products, "orders": orders,
            "formula": "Risk = min(100, 18 × active disruptions + 55 × shortage ratio + 12 × impacted suppliers). This is a transparent prototype heuristic, not an industry-certified risk model."}

@app.get("/api/dashboard")
@auth_required
def dashboard():
    result = build_analysis(request.current_user)
    result["disruptions"] = result["disruptions"][:5]
    result["recent_orders"] = result["orders"][:5]
    return jsonify(result)

@app.post("/api/analysis/run")
@auth_required
def run_analysis():
    result = build_analysis(request.current_user)
    saved = insert_doc("analyses", {"owner_id": request.current_user["id"], "workspace": workspace_of(request.current_user), "risk_score": result["risk_score"], "risk_level": result["risk_level"],
                                    "shortage": result["shortage"], "formula": result["formula"], "created_at": now_iso()})
    notify_workspace(request.current_user, "Impact analysis completed", f"Network risk score: {result['risk_score']}/100 ({result['risk_level']}).", "warning" if result["risk_score"] >= 35 else "success")
    result["analysis_id"] = saved["id"]
    return jsonify(result)

@app.post("/api/recovery/plan")
@auth_required
def recovery_plan():
    data = request.get_json(silent=True) or {}
    shortage = max(0, float(data.get("shortage", 0) or 0))
    budget = max(0, float(data.get("budget", 0) or 0))
    deadline = max(0, int(data.get("deadline_days", 7) or 7))
    suppliers = find_docs("suppliers", owner_query(request.current_user))
    candidates = []
    for s in suppliers:
        cap = number(s, "capacity")
        cost = number(s, "unit_cost")
        lead = int(number(s, "lead_days", 1))
        reliability = number(s, "reliability", 95)
        if cap <= 0 or cost < 0: continue
        qty = min(shortage, cap)
        total_cost = qty * cost
        feasible = lead <= deadline and total_cost <= budget and qty > 0
        candidates.append({"supplier": s.get("name"), "supplier_id": s.get("id"), "capacity": cap, "unit_cost": cost, "lead_days": lead,
                           "reliability": reliability, "quantity": qty, "total_cost": round(total_cost,2), "feasible": feasible,
                           "reason": "Meets capacity, budget and deadline constraints" if feasible else "; ".join([x for x, ok in [("delivery exceeds deadline", lead <= deadline), ("cost exceeds budget", total_cost <= budget), ("no usable capacity", qty > 0)] if not ok])})
    candidates.sort(key=lambda x: (not x["feasible"], x["total_cost"], x["lead_days"], -x["reliability"]))
    selected = next((x for x in candidates if x["feasible"] and x["quantity"] >= shortage), None)
    if not selected:
        selected = next((x for x in candidates if x["feasible"]), None)
    remaining = max(0, shortage - (selected["quantity"] if selected else 0))
    result = {"requested_shortage": shortage, "budget": budget, "deadline_days": deadline, "candidates": candidates,
              "selected": selected, "remaining_shortage": remaining, "status": "Feasible" if selected and remaining == 0 else "Partially feasible" if selected else "No feasible supplier found",
              "explanation": "Candidates are ranked by feasibility, cost, delivery time, and reliability. This baseline chooses one supplier; multi-supplier allocation can be added as an optimization extension."}
    insert_doc("analyses", {"owner_id": request.current_user["id"], "workspace": workspace_of(request.current_user), "kind": "recovery_plan", "result": result, "created_at": now_iso()})
    notify_workspace(request.current_user, "Recovery plan generated", result["status"], "success" if result["status"] == "Feasible" else "warning")
    return jsonify(result)

@app.get("/api/analyses/history")
@auth_required
def analyses_history():
    return jsonify({"items": find_docs("analyses", owner_query(request.current_user), 100)})

@app.get("/api/reports/export")
@auth_required
def export_report():
    kind = request.args.get("type", "pdf").lower()
    data = build_analysis(request.current_user)
    title = "RippleGuard Supply Chain Risk Report"
    lines = [title, f"Company: {request.current_user.get('company','')}", f"Generated (UTC): {now_iso()}", "", "SUMMARY",
             f"Suppliers: {data['supplier_count']}", f"Facilities: {data['facility_count']}", f"Products: {data['product_count']}",
             f"Orders: {data['order_count']}", f"Active disruptions: {data['active_disruptions']}", f"Estimated shortage: {data['shortage']}",
             f"Risk score: {data['risk_score']}/100 ({data['risk_level']})", "", "RISK FORMULA", data["formula"], "", "ACTIVE DISRUPTIONS"]
    for d in data["disruptions"]:
        lines.append(f"- {d.get('name','Incident')} | Supplier: {d.get('affected_supplier','Not specified')} | Status: {d.get('status','Active')} | Severity: {d.get('severity','Not specified')}")
    if kind == "csv":
        output = io.StringIO(); writer = csv.writer(output); writer.writerow(["Metric", "Value"])
        for line in lines: writer.writerow([line])
        return send_file(io.BytesIO(output.getvalue().encode("utf-8-sig")), mimetype="text/csv", as_attachment=True, download_name="rippleguard-report.csv")
    if kind == "docx":
        try:
            from docx import Document
            doc = Document(); doc.add_heading(title, 0)
            for line in lines[1:]:
                if line in ["SUMMARY", "RISK FORMULA", "ACTIVE DISRUPTIONS"]: doc.add_heading(line.title(), level=1)
                else: doc.add_paragraph(line)
            out = io.BytesIO(); doc.save(out); out.seek(0)
            return send_file(out, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document", as_attachment=True, download_name="rippleguard-report.docx")
        except Exception as exc: return jsonify({"error": f"DOCX export unavailable: {exc}"}), 500
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
        out = io.BytesIO(); doc = SimpleDocTemplate(out, pagesize=letter); styles = getSampleStyleSheet(); story = []
        for i, line in enumerate(lines):
            if i == 0: story.append(Paragraph(line, styles["Title"]))
            elif line in ["SUMMARY", "RISK FORMULA", "ACTIVE DISRUPTIONS"]: story.append(Spacer(1,8)); story.append(Paragraph(line, styles["Heading2"]))
            else: story.append(Paragraph(line.replace("&", "&amp;"), styles["BodyText"]))
            story.append(Spacer(1,4))
        doc.build(story); out.seek(0)
        return send_file(out, mimetype="application/pdf", as_attachment=True, download_name="rippleguard-report.pdf")
    except Exception as exc:
        return jsonify({"error": f"PDF export unavailable: {exc}. Install reportlab."}), 500


@app.get("/api/orders/<order_id>/report")
@auth_required
def order_lifecycle_report(order_id):
    user = request.current_user
    order = next((x for x in find_docs("orders", owner_query(user), 1000) if x.get("id") == order_id), None)
    if not order:
        return jsonify({"error": "Order not found in this workspace."}), 404
    steps = [x for x in find_docs("process_steps", owner_query(user), 1000) if str(x.get("order_id")) == order_id]
    steps.sort(key=lambda x: x.get("started_at", ""))
    kind = request.args.get("type", "pdf").lower()
    lines = ["RippleGuard Product Journey Report", f"Order: {order.get('order_no') or order.get('name','')}", f"Product: {order.get('name','')}",
             f"Quantity: {order.get('quantity',0)}", f"Status: {order.get('status','Open')}", f"Supplier: {order.get('supplier','Not specified')}",
             f"Start: {order.get('start_date','Not recorded')}", f"Due date: {order.get('due_date','Not recorded')}", f"Generated (UTC): {now_iso()}", "", "PROCESS TIMELINE"]
    for i, step in enumerate(steps, 1):
        lines.append(f"{i}. {step.get('stage','Process step')} | Supplier: {step.get('supplier','Not specified')} | Location: {step.get('location','Not specified')} | Start: {step.get('started_at','')} | Complete: {step.get('completed_at','Not completed')} | Status: {step.get('status','In progress')} | Notes: {step.get('notes','')}")
    if kind == "csv":
        output = io.StringIO(); writer = csv.writer(output); writer.writerow(["Field", "Value"])
        for line in lines: writer.writerow([line])
        return send_file(io.BytesIO(output.getvalue().encode("utf-8-sig")), mimetype="text/csv", as_attachment=True, download_name="rippleguard-product-journey.csv")
    if kind == "docx":
        from docx import Document
        doc = Document(); doc.add_heading(lines[0], 0)
        for line in lines[1:]:
            if line == "PROCESS TIMELINE": doc.add_heading(line.title(), level=1)
            else: doc.add_paragraph(line)
        out = io.BytesIO(); doc.save(out); out.seek(0)
        return send_file(out, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document", as_attachment=True, download_name="rippleguard-product-journey.docx")
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    out = io.BytesIO(); doc = SimpleDocTemplate(out, pagesize=letter); styles = getSampleStyleSheet(); story = []
    for i, line in enumerate(lines):
        if i == 0: story.append(Paragraph(line, styles["Title"]))
        elif line == "PROCESS TIMELINE": story.append(Spacer(1,8)); story.append(Paragraph(line.title(), styles["Heading2"]))
        else: story.append(Paragraph(line.replace("&", "&amp;"), styles["BodyText"]))
        story.append(Spacer(1,4))
    doc.build(story); out.seek(0)
    return send_file(out, mimetype="application/pdf", as_attachment=True, download_name="rippleguard-product-journey.pdf")

@app.errorhandler(404)
def not_found(_): return jsonify({"error": "API route not found."}), 404


@app.post("/api/demo/seed")
@auth_required
def seed_demo():
    user = request.current_user
    existing = any(find_docs(name, owner_query(user), 1) for name in ["suppliers", "facilities", "products", "orders", "disruptions"])
    if existing:
        return jsonify({"error": "This workspace already contains data. Demo records are only added to an empty workspace."}), 409
    samples = {
        "suppliers": [
            {"name":"Aster Components","category":"Electronic components","contact":"Operations Desk","capacity":120,"unit_cost":18,"lead_days":2,"reliability":96,"status":"Active","latitude":13.0827,"longitude":80.2707},
            {"name":"Northstar Materials","category":"Raw materials","contact":"Supply Team","capacity":80,"unit_cost":15,"lead_days":3,"reliability":91,"status":"At risk","latitude":12.9716,"longitude":77.5946},
            {"name":"Cobalt Logistics Supply","category":"Packaging components","contact":"Dispatch Team","capacity":60,"unit_cost":12,"lead_days":1,"reliability":98,"status":"Active","latitude":11.0168,"longitude":76.9558}
        ],
        "facilities": [
            {"name":"Coimbatore Assembly Plant","type":"Factory","address":"Coimbatore, Tamil Nadu","latitude":11.0168,"longitude":76.9558,"capacity":300,"status":"Operational"},
            {"name":"Bengaluru Distribution Hub","type":"Distribution hub","address":"Bengaluru, Karnataka","latitude":12.9716,"longitude":77.5946,"capacity":450,"status":"Operational"},
            {"name":"Chennai Receiving Warehouse","type":"Warehouse","address":"Chennai, Tamil Nadu","latitude":13.0827,"longitude":80.2707,"capacity":500,"status":"At risk"}
        ],
        "products": [
            {"name":"Control Module A","sku":"CM-A-001","category":"Electronics","required_components":160,"available_stock":95,"priority":"High","unit":"pieces","target_date":"2026-10-20"},
            {"name":"Smart Sensor Kit","sku":"SSK-210","category":"IoT components","required_components":90,"available_stock":70,"priority":"Critical","unit":"kits","target_date":"2026-10-18"},
            {"name":"Protective Enclosure","sku":"PE-014","category":"Packaging","required_components":120,"available_stock":130,"priority":"Medium","unit":"pieces","target_date":"2026-10-24"}
        ],
        "orders": [
            {"order_no":"RG-ORD-1001","name":"Control Module A batch","quantity":160,"status":"At risk","due_date":"2026-10-20","priority":"High"},
            {"order_no":"RG-ORD-1002","name":"Smart Sensor Kit batch","quantity":90,"status":"Open","due_date":"2026-10-18","priority":"Critical"}
        ],
        "disruptions": [
            {"name":"Northstar delivery delay","type":"Transport delay","affected_supplier":"Northstar Materials","location":"Bengaluru, Karnataka","severity":"High","status":"Active","start_date":"2026-10-14T09:00","expected_days":3,"description":"Illustrative training scenario: inbound raw material delivery is delayed."}
        ]
    }
    for name, rows in samples.items():
        for row in rows:
            row["owner_id"] = user["id"]; row["workspace"] = workspace_of(user)
            insert_doc(name, row)
    notify_workspace(user, "Demo workspace loaded", "Illustrative suppliers, facilities, products, orders and one disruption were added.", "success")
    return jsonify({"message":"Illustrative demo data added.","counts":{k:len(v) for k,v in samples.items()}}),201

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "5100")), debug=os.getenv("FLASK_DEBUG", "false").lower() == "true")
